"""The audit workbook: layout, writing and reading.

Columns are found by header name when reading, so a workbook that has been
round-tripped through Google Sheets (or had columns reordered) still reads.
"""

import datetime

METHOD_VERSION = "1.0"

PLATFORMS = ["Spotify", "YouTube Music", "Apple Music", "Musixmatch", "Genius", "Google top result"]
NO_LYRICS = "No lyrics found"

LYRICS = ["Yes", "Partial", "No"]
TONE_MARKS = ["Full", "Some", "None", "N/A"]
YES_NO_NA = ["Yes", "No", "N/A"]
COMMENT_STATUS = ["Live", "Held for review", "Removed", "Pinned", "Replied to by creator"]

SONG_COLUMNS = ["song_id", "title", "artist", "year", "era", "genre", "accuracy_rank",
                "frame_source", "url"]
PRESENCE_COLUMNS = ["song_id", "title", "artist", "platform", "lyrics", "tone_marks",
                    "translation", "synced", "published_by", "page_url", "checked_on", "notes"]
ACCURACY_COLUMNS = ["song_id", "title", "genre", "accuracy_rank", "line_slot",
                    "source_platform", "line_as_found", "correct_line", "checked_by", "notes"]
POST_COLUMNS = ["post_no", "link_tag", "video_url", "channel", "video_title",
                "views_at_post", "posted_on", "comment_url", "comment_status", "notes"]
DAILY_COLUMNS = ["date", "post_no", "comment_likes", "replies", "notes"]
SIGNUP_COLUMNS = ["signed_up_at", "link_tag", "answer"]

INPUT_FILL = "FFF2CC"  # pale yellow: fill these in
FIXED_FILL = "EDEDED"  # grey: written by the sampler, do not edit
HEADER_FILL = "3F3A34"
FONT = "Arial"


def era(year):
    try:
        y = int(str(year).strip()[:4])
    except ValueError:
        return "unknown"
    if y < 1980:
        return "pre-1980"
    if y < 2000:
        return "1980-1999"
    if y < 2015:
        return "2000-2014"
    return "2015-present"


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("The audit commands need openpyxl: pip install openpyxl") from exc
    return openpyxl


def write_workbook(path, songs, meta, accuracy_rows_per_genre, lines_per_song=5,
                   demand_posts=10, demand_days=14):
    openpyxl = _openpyxl()
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    base = Font(name=FONT, size=10)
    bold = Font(name=FONT, size=10, bold=True)
    header_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    input_fill = PatternFill("solid", fgColor=INPUT_FILL)
    fixed_fill = PatternFill("solid", fgColor=FIXED_FILL)

    def table(ws, columns, rows, inputs, widths):
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([row.get(c) for c in columns])
        last = max(ws.max_row, 2)
        for col_idx, name in enumerate(columns, start=1):
            letter = get_column_letter(col_idx)
            ws.column_dimensions[letter].width = widths.get(name, 14)
            fill = input_fill if name in inputs else fixed_fill
            for r in range(2, last + 1):
                cell = ws.cell(row=r, column=col_idx)
                cell.font = base
                if rows:
                    cell.fill = fill
        ws.freeze_panes = "A2"
        return last

    def dropdown(ws, columns, name, options, first, last):
        letter = get_column_letter(columns.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list",
                            error="Choose one of: " + ", ".join(options))
        dv.add(f"{letter}{first}:{letter}{last}")
        ws.add_data_validation(dv)

    # --- Read me --------------------------------------------------------------
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    readme = [
        ("Ogbifọ audit workbook", None),
        ("Method version", METHOD_VERSION),
        ("Sample drawn", meta["created_on"]),
        ("Random seed", str(meta["seed"])),
        ("Songs per genre", str(meta["per_genre"])),
        ("Accuracy songs per genre", str(meta["accuracy_per_genre"])),
        ("Candidate frame", meta["frame"]),
        ("", None),
        ("Legend", None),
        ("Yellow cells", "Fill these in. Use the dropdowns where there is one."),
        ("Grey cells", "Written by the sampler. Do not edit or re-sort them; the report relies on them."),
        ("", None),
        ("Steps", None),
        ("1. Presence", "For every row, open the platform, search for the song and record what a "
                        "listener would find. One row per song per platform."),
        ("2. Accuracy", f"Work through songs in accuracy_rank order within each genre until "
                        f"{meta['accuracy_per_genre']} per genre are checked. Skip a song only if "
                        f"no lyrics exist anywhere: set source_platform to '{NO_LYRICS}' on its "
                        "first line. Never skip a song for being hard."),
        ("", "Source: use the platform with the most tone marks for that song (ties: the order "
             "of the Presence tab). Lines: the first 5 mainly-Yorùbá lines, starting at the first "
             "verse. Paste line_as_found exactly, errors included. A qualified teacher writes "
             "correct_line with spoken (lexical) tones."),
        ("3. Demand", "Log each comment in Demand posts, give each a distinct link_tag in its "
                      "waitlist link, record likes and replies daily, and paste form responses "
                      "into Signups."),
        ("4. Report", "python -m ogbifo_bench audit report <this file> --out ../audit/runs/report.md"),
        ("", None),
        ("Definitions", None),
        ("lyrics", "Yes: at least 80% of the song. Partial: a snippet or some verses. No: none."),
        ("tone_marks", "Full: nearly every word that needs a tone mark has one (right or wrong; "
                       "correctness is measured in Accuracy). Some: marks on some words only. "
                       "None: no tone marks, even if underdots are present. N/A: no lyrics."),
        ("translation", "Yes if an English or other translation is on the same page."),
        ("synced", "Yes if lines are time-synced to the audio. N/A where the platform has no sync."),
        ("published_by", "Who supplied the lyrics, e.g. Musixmatch, LyricFind, a fan blog, the artist."),
        ("", None),
        ("Example rows", None),
        ("Presence", "song_id kwam1-example · Musixmatch · lyrics Partial · tone_marks Some · "
                     "translation No · synced Yes · published_by Musixmatch community · "
                     "checked_on 2026-09-29"),
        ("Accuracy", "source_platform Genius · line_as_found 'Oko mi wa nile' · correct_line "
                     "'Ọkọ mi wà nílé' · checked_by Adé (teacher)"),
    ]
    for label, value in readme:
        ws.append([label, value])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = base
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, (label, value) in enumerate(readme, start=1):
        if value is None and label:
            ws.cell(row=r, column=1).font = Font(name=FONT, size=12 if r == 1 else 10, bold=True)
    ws.cell(row=10, column=2).fill = input_fill
    ws.cell(row=11, column=2).fill = fixed_fill

    # --- Songs ----------------------------------------------------------------
    ws = wb.create_sheet("Songs")
    song_cols = SONG_COLUMNS + ["platforms_checked", "lyrics_anywhere", "tone_marks_anywhere"]
    last_song = table(ws, SONG_COLUMNS, songs, inputs=(),
                      widths={"title": 30, "artist": 24, "frame_source": 24, "url": 30})
    presence_last = 1 + len(songs) * len(PLATFORMS)
    p_id = f"Presence!$A$2:$A${presence_last}"
    p_lyrics = f"Presence!$E$2:$E${presence_last}"
    p_tones = f"Presence!$F$2:$F${presence_last}"
    first_calc = len(SONG_COLUMNS) + 1
    for offset, name in enumerate(song_cols[len(SONG_COLUMNS):]):
        cell = ws.cell(row=1, column=first_calc + offset, value=name)
        cell.font, cell.fill = header_font, header_fill
        ws.column_dimensions[get_column_letter(first_calc + offset)].width = 18
    checked_col = get_column_letter(first_calc)
    for r in range(2, len(songs) + 2):
        checked = f"{checked_col}{r}"
        formulas = [
            f'=COUNTIFS({p_id},$A{r},{p_lyrics},"?*")',
            f'=IF({checked}<{len(PLATFORMS)},"",IF(COUNTIFS({p_id},$A{r},{p_lyrics},"Yes")'
            f'+COUNTIFS({p_id},$A{r},{p_lyrics},"Partial")>0,"Yes","No"))',
            f'=IF({checked}<{len(PLATFORMS)},"",IF(COUNTIFS({p_id},$A{r},{p_lyrics},"Yes",{p_tones},"Full")'
            f'+COUNTIFS({p_id},$A{r},{p_lyrics},"Yes",{p_tones},"Some")'
            f'+COUNTIFS({p_id},$A{r},{p_lyrics},"Partial",{p_tones},"Full")'
            f'+COUNTIFS({p_id},$A{r},{p_lyrics},"Partial",{p_tones},"Some")>0,"Yes","No"))',
        ]
        for offset, formula in enumerate(formulas):
            ws.cell(row=r, column=first_calc + offset, value=formula).font = base
    ws.cell(row=1, column=first_calc).comment = Comment(
        "Calculated from the Presence tab. A song counts as fully checked at "
        f"{len(PLATFORMS)}; the two columns after this stay blank until then.", "Ogbifọ")

    # --- Presence -------------------------------------------------------------
    ws = wb.create_sheet("Presence")
    presence = [{"song_id": s["song_id"], "title": s["title"], "artist": s["artist"],
                 "platform": p} for s in songs for p in PLATFORMS]
    last = table(ws, PRESENCE_COLUMNS, presence,
                 inputs=set(PRESENCE_COLUMNS[4:]),
                 widths={"title": 28, "artist": 22, "platform": 18, "published_by": 22,
                         "page_url": 30, "notes": 30})
    dropdown(ws, PRESENCE_COLUMNS, "lyrics", LYRICS, 2, last)
    dropdown(ws, PRESENCE_COLUMNS, "tone_marks", TONE_MARKS, 2, last)
    dropdown(ws, PRESENCE_COLUMNS, "translation", YES_NO_NA, 2, last)
    dropdown(ws, PRESENCE_COLUMNS, "synced", YES_NO_NA, 2, last)
    for r in range(2, last + 1):
        ws.cell(row=r, column=PRESENCE_COLUMNS.index("checked_on") + 1).number_format = "yyyy-mm-dd"

    # --- Accuracy -------------------------------------------------------------
    ws = wb.create_sheet("Accuracy")
    accuracy = [
        {"song_id": s["song_id"], "title": s["title"], "genre": s["genre"],
         "accuracy_rank": s["accuracy_rank"], "line_slot": slot}
        for s in songs if s["accuracy_rank"] <= accuracy_rows_per_genre
        for slot in range(1, lines_per_song + 1)
    ]
    accuracy.sort(key=lambda a: (a["genre"], a["accuracy_rank"], a["line_slot"]))
    last = table(ws, ACCURACY_COLUMNS, accuracy, inputs=set(ACCURACY_COLUMNS[5:]),
                 widths={"title": 28, "line_as_found": 45, "correct_line": 45, "notes": 30})
    dropdown(ws, ACCURACY_COLUMNS, "source_platform", PLATFORMS + [NO_LYRICS], 2, last)
    for name in ("line_as_found", "correct_line"):
        ws.column_dimensions[get_column_letter(ACCURACY_COLUMNS.index(name) + 1)].width = 45

    # --- Demand ---------------------------------------------------------------
    ws = wb.create_sheet("Demand posts")
    posts = [{"post_no": n, "link_tag": f"v{n:02d}"} for n in range(1, demand_posts + 1)]
    last = table(ws, POST_COLUMNS, posts, inputs=set(POST_COLUMNS[2:]),
                 widths={"video_url": 30, "video_title": 30, "comment_url": 30, "notes": 30})
    dropdown(ws, POST_COLUMNS, "comment_status", COMMENT_STATUS, 2, last)
    for r in range(2, last + 1):
        ws.cell(row=r, column=POST_COLUMNS.index("posted_on") + 1).number_format = "yyyy-mm-dd"

    ws = wb.create_sheet("Demand daily")
    daily_rows = demand_posts * demand_days
    table(ws, DAILY_COLUMNS, [{} for _ in range(daily_rows)], inputs=set(DAILY_COLUMNS),
          widths={"notes": 30})
    for r in range(2, daily_rows + 2):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"

    ws = wb.create_sheet("Signups")
    table(ws, SIGNUP_COLUMNS, [], inputs=set(SIGNUP_COLUMNS),
          widths={"signed_up_at": 20, "answer": 50})

    # --- Dashboard ------------------------------------------------------------
    ws = wb.create_sheet("Dashboard", 1)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    n = len(songs)

    def songs_rng(col):
        return f"Songs!${col}$2:${col}${last_song}"

    def calc(offset):
        return get_column_letter(first_calc + offset)

    acc_last = 1 + len(accuracy)
    signup_rows = 5000
    rows = [
        ("Progress", None, None),
        ("Songs in sample", f"=COUNTA({songs_rng('A')})", None),
        ("Presence rows filled", f'=COUNTIF(Presence!$E$2:$E${presence_last},"?*")',
         f"=IF({n * len(PLATFORMS)}=0,\"\",B3/{n * len(PLATFORMS)})"),
        ("Songs fully checked", f"=COUNTIF({songs_rng(calc(0))},{len(PLATFORMS)})",
         "=IF(B2=0,\"\",B4/B2)"),
        ("Accuracy lines filled", f'=COUNTIF(Accuracy!$H$2:$H${acc_last},"?*")', None),
        ("", None, None),
        ("Headline (fully checked songs; the report adds confidence intervals)", "Songs", "Share"),
        ("Lyrics anywhere online", f'=COUNTIF({songs_rng(calc(1))},"Yes")',
         "=IF($B$4=0,\"\",B8/$B$4)"),
        ("Any tone marks anywhere", f'=COUNTIF({songs_rng(calc(2))},"Yes")',
         "=IF($B$4=0,\"\",B9/$B$4)"),
        ("", None, None),
        ("Lyrics found, by platform (checked rows)", "Found", "Share"),
    ]
    for p in PLATFORMS:
        r = len(rows) + 1
        platform_rng = f"Presence!$D$2:$D${presence_last}"
        lyrics_rng = f"Presence!$E$2:$E${presence_last}"
        rows.append((
            p,
            f'=COUNTIFS({platform_rng},"{p}",{lyrics_rng},"Yes")'
            f'+COUNTIFS({platform_rng},"{p}",{lyrics_rng},"Partial")',
            f'=IF(COUNTIFS({platform_rng},"{p}",{lyrics_rng},"?*")=0,"",'
            f'B{r}/COUNTIFS({platform_rng},"{p}",{lyrics_rng},"?*"))',
        ))
    rows.append(("", None, None))
    rows.append(("Demand test", None, None))
    signups_row = len(rows) + 1
    rows.append(("Waitlist signups", f"=COUNTA(Signups!$A$2:$A${signup_rows})", None))
    rows.append(("Decision gate",
                 f'=IF(B{signups_row}>=50,"Continue as planned",IF(B{signups_row}>=20,'
                 f'"Repeat on 10 more videos","Lead with tone-marker and captions"))', None))
    for label, b, c in rows:
        ws.append([label, b, c])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = base
    for r, (label, b, c) in enumerate(rows, start=1):
        if b is None or not str(b).startswith("="):
            for col in range(1, 4):
                ws.cell(row=r, column=col).font = bold
        if c and str(c).startswith("="):
            ws.cell(row=r, column=3).number_format = "0.0%"
    ws.cell(row=3, column=3).number_format = "0.0%"
    ws.cell(row=4, column=3).number_format = "0.0%"

    # --- Meta (read by the report) --------------------------------------------
    ws = wb.create_sheet("Meta")
    ws.sheet_state = "hidden"
    for key, value in meta.items():
        ws.append([key, str(value)])

    wb.save(path)


def _rows(ws):
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    header = [str(h).strip() if h is not None else "" for h in values[0]]
    out = []
    for row in values[1:]:
        record = {h: row[i] if i < len(row) else None for i, h in enumerate(header) if h}
        if any(v not in (None, "") for v in record.values()):
            out.append(record)
    return out


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.date().isoformat() if isinstance(value, datetime.datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_workbook(path):
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = {}
    for name, key in [("Songs", "songs"), ("Presence", "presence"), ("Accuracy", "accuracy"),
                      ("Demand posts", "posts"), ("Demand daily", "daily"),
                      ("Signups", "signups")]:
        if name not in wb.sheetnames:
            raise ValueError(f"{path}: missing the '{name}' tab")
        sheets[key] = [{k: _text(v) for k, v in row.items()} for row in _rows(wb[name])]
    meta = {}
    if "Meta" in wb.sheetnames:
        for row in wb["Meta"].iter_rows(values_only=True):
            if row and row[0]:
                meta[str(row[0])] = _text(row[1])
    sheets["meta"] = meta
    wb.close()
    return sheets
