import csv
import datetime
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from ogbifo_bench.audit import report
from ogbifo_bench.audit.sample import draw, load_candidates
from ogbifo_bench.audit.workbook import NO_LYRICS, PLATFORMS, era

try:
    import openpyxl
except ImportError:  # the audit needs it; the benchmark does not
    openpyxl = None

BENCH = Path(__file__).resolve().parent.parent
GENRES = ["Jùjú", "Fújì", "Highlife"]


def write_candidates(path, per_genre=12):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["title", "artist", "year", "genre"])
        for g in GENRES:
            for i in range(per_genre):
                w.writerow([f"Song {i}", f"{g} artist {i}", 1970 + i * 4, g])
        w.writerow(["Song 0", "Jùjú artist 0", 1970, "Jùjú"])  # duplicate listing


class StatsTest(unittest.TestCase):
    def test_wilson(self):
        low, high = report.wilson(5, 10)
        self.assertAlmostEqual(low, 0.2366, places=3)
        self.assertAlmostEqual(high, 0.7634, places=3)
        self.assertAlmostEqual(report.wilson(0, 10)[1], 0.2775, places=3)
        self.assertEqual(report.wilson(0, 0), (0.0, 0.0))

    def test_era(self):
        self.assertEqual([era(1979), era(1980), era("2014"), era(2015), era("")],
                         ["pre-1980", "1980-1999", "2000-2014", "2015-present", "unknown"])


class SampleTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "candidates.csv"
        write_candidates(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_duplicates_dropped(self):
        self.assertEqual(len(load_candidates(self.path)), 36)

    def test_draw_is_reproducible_and_stratified(self):
        candidates = load_candidates(self.path)
        first, _ = draw(candidates, 5, seed=7)
        again, _ = draw(candidates, 5, seed=7)
        other, _ = draw(candidates, 5, seed=8)
        self.assertEqual(first, again)
        self.assertNotEqual([s["song_id"] for s in first], [s["song_id"] for s in other])
        for g in GENRES:
            ranks = sorted(s["accuracy_rank"] for s in first if s["genre"] == g)
            self.assertEqual(ranks, [1, 2, 3, 4, 5])

    def test_shortfall_reported(self):
        _, shortfalls = draw(load_candidates(self.path), 20, seed=1)
        self.assertEqual(shortfalls, {"Fújì": 12, "Highlife": 12, "Jùjú": 12})


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class WorkbookRoundTripTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        tmp = Path(self.tmp.name)
        write_candidates(tmp / "candidates.csv")
        self.xlsx = tmp / "audit.xlsx"
        subprocess.run([sys.executable, "-m", "ogbifo_bench", "audit", "sample",
                        str(tmp / "candidates.csv"), "--out", str(self.xlsx),
                        "--per-genre", "10", "--accuracy-per-genre", "2", "--seed", "3"],
                       cwd=BENCH, check=True, capture_output=True)

    def tearDown(self):
        self.tmp.cleanup()

    def fill(self, skip_rank_one=True):
        wb = openpyxl.load_workbook(self.xlsx)
        songs = [r[0] for r in wb["Songs"].iter_rows(min_row=2, values_only=True)]
        checked = songs[:8]
        for row in wb["Presence"].iter_rows(min_row=2):
            song_id, platform = row[0].value, row[3].value
            if song_id in checked:
                i = checked.index(song_id)
                found = platform == "Genius" and i < 4
                row[4].value = "Yes" if found else "No"
                row[5].value = ("Full" if i < 2 else "None") if found else "N/A"
                row[6].value = "Yes" if found and i == 0 else "No"
                row[10].value = datetime.date(2026, 10, 1)
        for row in wb["Accuracy"].iter_rows(min_row=2):
            rank, slot = row[3].value, row[4].value
            if rank == 1 and skip_rank_one and slot == 1:
                row[5].value = NO_LYRICS
            if rank in (2, 3):
                row[5].value = "Genius"
                row[6].value = "Oko mi wa nile" if slot == 1 else "Ọkọ mi wà nílé"
                row[7].value = "Ọkọ mi wà nílé"
        posts = wb["Demand posts"]
        posts.cell(2, 3).value = "https://example.com/video"
        for i in range(55):
            wb["Signups"].append(["2026-10-02", "v01", "a song"])
        wb.save(self.xlsx)

    def test_report_numbers(self):
        from ogbifo_bench.audit.workbook import read_workbook
        self.fill()
        summary = report.build(read_workbook(self.xlsx))
        p, a, d = summary["presence"], summary["accuracy"], summary["demand"]
        self.assertEqual(p["songs_complete"], 8)
        self.assertEqual(p["headline"]["lyrics"]["count"], 4)
        self.assertEqual(p["headline"]["any_marks"]["count"], 2)
        self.assertEqual(p["headline"]["translation"]["count"], 1)
        self.assertEqual(p["platforms"]["Genius"]["lyrics"]["count"], 4)
        self.assertEqual(a["songs_checked"], 6)  # ranks 2 and 3 in each of 3 genres
        self.assertEqual(a["songs_without_lyrics"], 3)
        self.assertEqual(a["lines"], 30)
        self.assertEqual(a["lines_correct"]["count"], 24)
        self.assertEqual(a["warnings"], [])
        self.assertEqual(d["signups"], 55)
        self.assertEqual(d["gate"], "Continue as planned")
        markdown = report.to_markdown(summary)
        self.assertIn("4 of 8 Yorùbá songs we checked have lyrics", markdown)

    def test_skipped_rank_is_flagged(self):
        from ogbifo_bench.audit.workbook import read_workbook
        self.fill(skip_rank_one=False)
        warnings = report.build(read_workbook(self.xlsx))["accuracy"]["warnings"]
        self.assertEqual(len([w for w in warnings if "no longer random" in w]), 3)

    def test_presence_rows_cover_every_platform(self):
        wb = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(wb["Presence"].max_row - 1, 30 * len(PLATFORMS))

    def test_report_cli(self):
        self.fill()
        out = Path(self.tmp.name) / "report.md"
        subprocess.run([sys.executable, "-m", "ogbifo_bench", "audit", "report",
                        str(self.xlsx), "--out", str(out)],
                       cwd=BENCH, check=True, capture_output=True)
        self.assertTrue(out.with_suffix(".json").exists())


if __name__ == "__main__":
    unittest.main()
